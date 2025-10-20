from django.shortcuts import render
from rest_framework.response import Response
from .serializers import CustomerSerializer
from rest_framework import viewsets
from rest_framework.filters import OrderingFilter, SearchFilter
from .models import Customer
from rest_framework.permissions import IsAuthenticated
from django.contrib.auth.decorators import login_required
from rest_framework import status
import openpyxl
from openpyxl.utils import get_column_letter
from django.http import HttpResponse
from rest_framework.decorators import action
from django.core.cache import cache  # 引入cache
import json
import hashlib
from rest_framework.pagination import PageNumberPagination


class CustomerPagination(PageNumberPagination):
    page_size_query_param = 'page_size'
    max_page_size = 100


class CustomersSet(viewsets.ModelViewSet):
    serializer_class = CustomerSerializer
    pagination_class = CustomerPagination
    filter_backends = [SearchFilter, OrderingFilter]
    search_fields = ["name", "email", "phone", "address"]
    ordering_fields = ["created_at", "updated_at", "name"]
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return Customer.objects.filter(owner=self.request.user).select_related('owner')

    def perform_create(self, serializer):
        serializer.save(owner=self.request.user)

    def list(self, request, *args, **kwargs):
        """
        重写 list 方法，加入缓存逻辑
        """
        # 获取客户数据版本号
        version_key = 'crm:customers:version'
        version = cache.get(version_key, 1)  # 默认版本为1

        # 构造动态且唯一的缓存键
        query_params_str = json.dumps(sorted(request.query_params.items()))
        query_hash = hashlib.md5(query_params_str.encode('utf-8')).hexdigest()

        cache_key = f"crm:customers:list:u{request.user.id}:v{version}:{query_hash}"

        # 尝试从缓存获取数据
        cached_response_data = cache.get(cache_key)
        if cached_response_data:
            return Response(cached_response_data)

        # 缓存未命中，执行原始的查询和序列化逻辑
        response = super().list(request, *args, **kwargs)

        # 将返回结果的数据部分写入缓存
        cache.set(cache_key, response.data, timeout=3600)  # 设置TTL为1h
        print(f"写入缓存成功。Key: {cache_key}")

        return response

    @action(detail=False, methods=['post'])
    def import_excel(self, request):
        file = request.FILES.get("file")
        if not file:
            return Response({"error": "请上传文件"},
                            status=status.HTTP_400_BAD_REQUEST)
        try:
            wb = openpyxl.load_workbook(file)
            ws = wb.active
        except Exception as e:
            return Response({"error": f"文件解析失败：{str(e)}"},
                             status=status.HTTP_400_BAD_REQUEST)

        # 第一行是表头，从第二行开始
        created_count, updated_count = 0, 0
        for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row[1]:  # 必须有姓名
                continue
            customer_id, name, phone, email, address, owner_username, *_ = row

            # 如果有 ID，尝试更新；否则新建
            if customer_id:
                try:
                    customer = Customer.objects.get(id=customer_id)
                    customer.name = name
                    customer.phone = phone
                    customer.email = email
                    customer.address = address
                    customer.save()
                    updated_count += 1
                except Customer.DoesNotExist:
                    continue
            else:
                Customer.objects.create(
                    name=name,
                    phone=phone,
                    email=email,
                    address=address,
                    owner=request.user  # 默认当前用户
                )
                created_count += 1

        return Response({
            "msg": "导入完成",
            "created": created_count,
            "updated": updated_count
        })

    @action(detail=False, methods=['get'])
    def export_excel(self, request):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Customers"

        # 表头
        headers = ["ID", "姓名", "电话", "邮箱", "地址", "负责人", "创建时间", "更新时间"]
        ws.append(headers)

        # 数据行
        for c in self.get_queryset():
            ws.append([
                c.id,
                c.name,
                c.phone,
                c.email,
                c.address,
                c.owner.username if c.owner else "",
                c.created_at.strftime("%Y-%m-%d %H:%M:%S"),
                c.updated_at.strftime("%Y-%m-%d %H:%M:%S"),
            ])

        # 设置列宽
        for i, col in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = 20

        # 返回 response
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response['Content-Disposition'] = 'attachment; filename=customers.xlsx'
        wb.save(response)
        return response
